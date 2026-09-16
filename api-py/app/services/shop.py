"""店铺数据服务(JSON 表单 + Excel 上传 + 汇总;含 DoS/MIME 防护)。"""
import io
import json
import logging
import re

logger = logging.getLogger("api")
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook
from sqlalchemy import select, text
from sqlalchemy.orm import Session

from app.core.exceptions import ApiException
from app.db.models.shop_health_score import ShopHealthScore
from app.db.models.shop_product_count import ShopProductCount
from app.db.models.shop_product_data import ShopProductData
from app.db.models.shop_star_data import ShopStarData
from app.db.models.shop_trade_data import ShopTradeData
from app.db.models.shop_traffic_data import ShopTrafficData

# 仅声明 openpyxl 真正能解析的 OOXML 格式(.xlsx);旧版 .xls(BIFF/OLE2) openpyxl 无法解析,
# 不声明 application/vnd.ms-excel,避免"提示支持但必然失败"的路径(真正支持需引入 xlrd 等新依赖,按边界不新增)。
EXCEL_MIME = ["application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"]
MAX_SHEETS = 10
MAX_ROWS = 5000


class NumberParseError(ValueError):
    """数值解析失败:非空值无法解析为数值(由 parse_number 抛出)。"""


def parse_number(v: Any) -> Optional[Decimal]:
    """数值解析唯一 owner:库侧取值(DECIMAL/str)与 Excel 单元格共用同一套解析规则。

    职责边界(调用点不得再自行做类型判断或另造解析):
    - 本函数只回答"是不是数值、数值是多少",不决定解析失败后怎么办;
    - None / 空白 -> None(无值,不算脏值);
    - int/float/Decimal/bool 与字符串统一归一到 Decimal(经 str 中转,避免 float 二进制尾差);
    - 字符串先剥离千分位(, 与，)、百分号(% 与 ％)与首尾空白;
    - 其余(含无法解析的字符串)抛 NumberParseError,由调用方按各自策略处理:
        _to_num(Excel 导入):记日志并跳过该单元格;
        _dec(读路径输出):输出 None,不让单条脏值把 /api/shop/summary 打成 500。
    """
    if v is None:
        return None
    if isinstance(v, Decimal):
        return v
    if isinstance(v, bool):  # bool 是 int 子类:显式按 0/1 处理(与既有 _to_num 行为一致)
        return Decimal(int(v))
    if isinstance(v, int):
        return Decimal(v)
    if isinstance(v, float):
        return Decimal(str(v))
    if isinstance(v, str):
        cleaned = v.replace("％", "").replace("%", "").replace("，", "").replace(",", "").strip()
        if cleaned == "":
            return None
        try:
            return Decimal(cleaned)
        except InvalidOperation as e:
            raise NumberParseError(f"无法解析为数值: {v!r}") from e
    raise NumberParseError(f"不支持的数值类型: {type(v).__name__}")


class DataDateParseError(ValueError):
    """日期解析失败:无法识别为可确定的日期(由 parse_data_date 抛出)。

    `reason` 是结构化英文枚举(《日期与数值校验方案》§8 口径),供调用方决定失败策略,
    并供 #PB-24-2 的逐行问题清单直接消费(单一真源 = 下方 DATE_PARSE_REASONS)。
    """

    def __init__(self, reason: str, message: str = "") -> None:
        self.reason = reason
        super().__init__(message or reason)


# 日期解析失败原因(结构化枚举;集中定义,禁止在调用点各写一套字面量)
DATE_PARSE_REASONS = (
    "missing_date",                      # 空 / 全空白 / 日期列命中无数据标记
    "ambiguous_date_format",             # 05/08/2026、05-08-2026、5/8/2026(月日顺序不可判定)
    "two_digit_year",                    # 26-08-05、26/08/05、26.08.05(两位年份,世纪歧义)
    "missing_year",                      # 8/5、08-05(无年份,不得默认当前年)
    "incomplete_date",                   # 2026-8、2026/08(缺日,不得默认月初)
    "invalid_calendar_date",             # 2026-02-30、2026-13-45(形态合法但非真实日期)
    "date_serial_number_not_supported",  # Excel 序列号(常规格式数字,如 45000)
    "date_out_of_range",                 # < 2000-01-01 或 > 服务器当天 + 1 天
    "unparseable_date",                  # 其余(含中文形态 2026年8月5日)
)

# 日期可接受范围(用户 2026-09-15 批准):下限固定 2000-01-01(项目 2026 上线,更早必为脏值);
# 上限 = 服务器当天 + 1 天 —— **时区容差**:表单默认值由客户端本地时区生成(DataForm.vue),
# 实测服务器本地为 UTC+8,客户端时区领先/设备时钟偏差时严格的「<= 今天」会误杀正常提交。
MIN_DATA_DATE = date(2000, 1, 1)
MAX_DATA_DATE_TOLERANCE_DAYS = 1

# 表单日期校验失败的用户可见文案(**单一真源**;口语化专句,不拼接通用兜底句):
# 日期是高频且用户能自行修正的错误,专句要回答"该怎么填"(#PB-24-1-R1 用户裁决)。
DATA_DATE_INVALID_MESSAGE = "数据日期不正确，请填写如 2026-08-05 这样的日期"

# 无数据标记(商智导出对「该指标本期无数据」的写法;真实样本实测 spu动销率='-')。
# 语义 = **NULL(无值)**,不是坏值:数值列命中 -> 写 NULL 且不跳过整行;日期列命中 -> missing_date。
# 白名单式:未列出的任何形态一律仍按坏值处理,避免把真脏值吞成 NULL。
NO_DATA_MARKERS = ("-", "—", "–", "−", "－", "--", "/", "\\", "N/A", "NA", "n/a")

# 日期白名单(唯一 owner;只认「年在前、无歧义」的形态,白名单外一律失败,不做任何猜测):
#   W1 datetime/date(Excel 真日期单元格);
#   W2 YYYY-MM-DD、W3 YYYY/M/D、W4 YYYY.M.D、W5 YYYYMMDD(恰好 8 位;str 或 int / 整值 float);
#   W6 上述形态 + 时间后缀(空格或 T 分隔,只取日期部分);
#   W7 区间 `起~止` -> 取【结束日】(= 数据截止日;用户 2026-09-15 裁决;当日导出首尾相同,故与旧行为一致)。
_DATA_DATE_SEPARATED = (
    re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})$"),
    re.compile(r"^(\d{4})/(\d{1,2})/(\d{1,2})$"),
    re.compile(r"^(\d{4})\.(\d{1,2})\.(\d{1,2})$"),
)
_DATA_DATE_COMPACT = re.compile(r"^(\d{4})(\d{2})(\d{2})$")
DATA_DATE_RANGE_SEPARATOR = "~"
_DATA_DATE_TIME_SEPARATORS = ("T", " ")
# 黑名单判据(只用于给出**结构化失败原因**;不改变「白名单外一律失败」的结论,也不做任何猜测)
_AMBIGUOUS_DATE = re.compile(r"^\d{1,2}[-/.]\d{1,2}[-/.]\d{4}$")       # 月/日在前
_TWO_DIGIT_YEAR_DATE = re.compile(r"^\d{2}[-/.]\d{1,2}[-/.]\d{1,2}$")  # 两位年份
_MISSING_YEAR_DATE = re.compile(r"^\d{1,2}[-/.]\d{1,2}$")              # 无年份
_INCOMPLETE_DATE = re.compile(r"^\d{4}[-/.]\d{1,2}$")                  # 缺日
_DIGITS_ONLY = re.compile(r"^\d+$")


def _assert_data_date_range(parsed: date) -> date:
    """范围校验(MIN_DATA_DATE ~ 服务器当天 + MAX_DATA_DATE_TOLERANCE_DAYS)并返回 date。"""
    if parsed < MIN_DATA_DATE or parsed > date.today() + timedelta(days=MAX_DATA_DATE_TOLERANCE_DAYS):
        raise DataDateParseError("date_out_of_range", f"日期超出允许范围: {parsed.isoformat()}")
    return parsed


def _validated_data_date(year: int, month: int, day: int) -> date:
    """真实日历校验(不依赖 MySQL 的宽松解析)+ 范围校验 -> date。"""
    try:
        parsed = date(year, month, day)
    except ValueError as e:
        raise DataDateParseError(
            "invalid_calendar_date", f"不是真实存在的日期: {year:04d}-{month:02d}-{day:02d}"
        ) from e
    return _assert_data_date_range(parsed)


def _parse_single_data_date_obj(value: Any) -> date:
    """解析单个日期(W1/W2-W6)为 date;失败 -> DataDateParseError(reason)。内部函数,调用方用 parse_data_date/data_date_span。"""
    if value is None:
        raise DataDateParseError("missing_date", "日期为空")
    if isinstance(value, datetime):   # datetime 是 date 的子类,必须先判 datetime
        return _assert_data_date_range(value.date())
    if isinstance(value, date):
        return _assert_data_date_range(value)
    if isinstance(value, bool):       # bool 是 int 子类:不是日期
        raise DataDateParseError("date_serial_number_not_supported", f"不支持的日期类型: {type(value).__name__}")
    if isinstance(value, int):
        text = str(value)
    elif isinstance(value, float):
        if not value.is_integer():
            raise DataDateParseError("date_serial_number_not_supported", f"不支持的日期数值: {value!r}")
        text = str(int(value))
    elif isinstance(value, str):
        text = value.strip()
    else:
        raise DataDateParseError("unparseable_date", f"不支持的日期类型: {type(value).__name__}")
    if text == "" or text in NO_DATA_MARKERS:
        raise DataDateParseError("missing_date", f"日期为空或无数据标记: {value!r}")
    raw_text = text          # 剥时间后缀前的原样文本:序列号判据只看它(避免把「8 位紧凑 + 时间」误报成序列号)
    # W5 必须「恰好 8 位」:先按原样匹配紧凑形态(因此 `20260805 13:45:00` 这类带后缀的紧凑串**不接受**,
    # 与方案 §4.1.1/§8.4 的口径一致);再剥时间后缀走 W6 + W2/W3/W4。
    compact = _DATA_DATE_COMPACT.match(text)
    if compact:
        return _validated_data_date(int(compact.group(1)), int(compact.group(2)), int(compact.group(3)))
    for separator in _DATA_DATE_TIME_SEPARATORS:   # W6:带时间后缀 -> 只取日期部分
        if separator in text:
            text = text.split(separator, 1)[0].strip()
            break
    for pattern in _DATA_DATE_SEPARATED:
        matched = pattern.match(text)
        if matched:
            return _validated_data_date(int(matched.group(1)), int(matched.group(2)), int(matched.group(3)))
    # 白名单未命中:按黑名单判据给出结构化原因(结论仍是失败,无任何猜测分支)
    if _AMBIGUOUS_DATE.match(text):
        raise DataDateParseError("ambiguous_date_format", f"日期月日顺序不可判定: {value!r}")
    if _TWO_DIGIT_YEAR_DATE.match(text):
        raise DataDateParseError("two_digit_year", f"日期使用了两位年份: {value!r}")
    if _MISSING_YEAR_DATE.match(text):
        raise DataDateParseError("missing_year", f"日期缺少年份: {value!r}")
    if _INCOMPLETE_DATE.match(text):
        raise DataDateParseError("incomplete_date", f"日期不完整: {value!r}")
    if _DIGITS_ONLY.match(raw_text):
        raise DataDateParseError("date_serial_number_not_supported", f"疑似 Excel 序列号: {value!r}")
    raise DataDateParseError("unparseable_date", f"无法识别为日期: {value!r}")


def _parse_single_data_date(value: Any) -> str:
    """单个日期的 ISO 形态(内部便捷包装)。"""
    return _parse_single_data_date_obj(value).isoformat()


def data_date_span(value: Any) -> Tuple[date, date]:
    """日期跨度(**与 parse_data_date 同一 owner**,不重复解析):

    - 区间 `起~止` -> (起, 止),两端都必须合法(判据与 parse_data_date 完全一致);
    - 单日形态(W1/W2-W6) -> (d, d);
    - 失败同样抛 DataDateParseError(reason),由调用方决定策略。

    用途:Excel 导入按**文件区间跨度**推导 `time_range`(见 time_range_from_span),避免把「当日导出」存进 7d 桶。
    """
    if isinstance(value, str) and DATA_DATE_RANGE_SEPARATOR in value:
        parts = value.split(DATA_DATE_RANGE_SEPARATOR)
        if len(parts) != 2:
            raise DataDateParseError("unparseable_date", f"区间日期无法识别: {value!r}")
        return _parse_single_data_date_obj(parts[0]), _parse_single_data_date_obj(parts[1])
    parsed = _parse_single_data_date_obj(value)
    return parsed, parsed


def parse_data_date(v: Any) -> Optional[str]:
    """日期解析唯一 owner(对称 parse_number):Excel 单元格与表单入参共用同一套规则。

    职责边界(调用点不得再自行做日期判断,也不得把归一权交回 MySQL):
    - 本函数只回答「是不是可识别的日期、规范化结果是什么」,**不决定**失败后是跳过还是 400;
    - 识别成功 -> 规范形态 YYYY-MM-DD(写库唯一形态);
    - 区间 `起~止` 取【结束日】(W7;两端都必须合法);
    - 其余一律抛 DataDateParseError(reason),reason 枚举见 DATE_PARSE_REASONS;
      调用方策略:表单 -> 400(DATA_DATE_INVALID_MESSAGE);(#PB-24-2 起)Excel -> 跳过该行 + 逐行清单。
    """
    _start, end = data_date_span(v)   # 复用同一 owner 的跨度解析(区间两端都会做合法性校验)
    return end.isoformat()            # W7:取结束日(数据截止日)


def _dec(v, scale):
    """读路径数值格式化:统一解析后按固定小数位输出(对齐 NestJS decimal 列输出的字符串形态)。
    解析失败(脏值)输出 None,不抛异常打断 /api/shop/summary。"""
    try:
        num = parse_number(v)
    except NumberParseError:
        logger.warning(f"_dec 跳过无法解析的数值: {v!r}")
        return None
    return f"{num:.{scale}f}" if num is not None else None


# DECIMAL 列序列化规格(唯一真源:project/scripts/schema.sql 的 DECIMAL 列 scale;
# 字段集合与 backend NestJS TypeORM decimal 列一一对应)。
# 用途:summary 输出把 DECIMAL 值统一字符串化(对齐 TypeORM decimal -> string,如 "1000.50"),
# 避免同一响应内 star 是字符串、trade 是数字的类型语义分裂。
DECIMAL_SCALES: Dict[str, Dict[str, int]] = {
    "star": {"shop_star": 1, "service_score": 1, "logistics_score": 1, "after_sale_score": 1, "product_score": 1},
    "trade": {"trade_amount": 2, "conversion_rate": 2, "customer_unit_price": 2,
              "avg_stay_duration": 2, "cart_conversion_rate": 2},
    "traffic": {"avg_stay_duration": 2, "product_avg_page_views": 2, "product_avg_stay_duration": 2,
                "uv_value": 2, "customer_unit_price": 2, "cart_conversion_rate": 2, "cart_amount": 2,
                "trade_conversion_rate": 2, "trade_amount": 2},
    "product": {"spu_active_rate": 2, "trade_amount": 2, "trade_conversion_rate": 2,
                "customer_unit_price": 2, "item_unit_price": 2, "spu_cart_rate": 2, "cart_amount": 2,
                "product_avg_page_views": 2, "product_avg_stay_duration": 2},
    "health_score": {"avg_score": 2},
}


def _serialize_decimals(payload: Dict[str, Any], table: str) -> Dict[str, Any]:
    """按表规格把 DECIMAL 字段统一序列化为字符串(固定小数位),其余字段原值透传。

    数值输出的唯一 owner:字段是否 decimal 由 DECIMAL_SCALES 决定,接口层不得再散落
    str()/float()/round() 包装。值为 None 时输出 null(与 NestJS 一致,不伪造 0)。
    """
    scales = DECIMAL_SCALES[table]
    return {k: (_dec(v, scales[k]) if k in scales else v) for k, v in payload.items()}


def _assert_phase2(db, merchant_id):
    from app.services.task_progress import get_phase2_unlock_state
    if not get_phase2_unlock_state(db, merchant_id)["unlocked"]:
        raise ApiException("阶段二未解锁", code=403, status_code=403)

def _upsert(db, model, values, conflict, commit: bool = True):
    """upsert;默认立即 commit(其它调用方行为不变)。upload_excel 传 commit=False 以整批一个事务。"""
    row = db.execute(select(model).where(*[getattr(model, k) == v for k, v in conflict.items()])).scalar_one_or_none()
    if row is None:
        db.add(model(**values))
    else:
        for k, v in values.items():
            setattr(row, k, v)
    if commit:
        db.commit()
    else:
        # 批内 flush:同一批次里出现相同 (merchant_id, data_date, time_range) 时,后续 SELECT 必须能看到前一行
        # 从而走 UPDATE(幂等覆盖)而不是 INSERT —— 否则撞唯一键 uk_merchant_date_range 打成 500。
        # 背景:SessionLocal 为 autoflush=False;旧实现整批不 flush,同一文件内的重复键必然 500(#PB-24-2 发现)。
        db.flush()

def upload_star(db, merchant_id, dto):
    _assert_phase2(db, merchant_id)
    data_date = _normalized_form_data_date(dto.get("dataDate"))
    _upsert(db, ShopStarData, {"merchant_id": merchant_id, "data_date": data_date, "shop_star": str(dto["shopStar"]),
        "service_score": str(dto["serviceScore"]) if dto.get("serviceScore") is not None else None,
        "logistics_score": str(dto["logisticsScore"]) if dto.get("logisticsScore") is not None else None,
        "after_sale_score": str(dto["afterSaleScore"]) if dto.get("afterSaleScore") is not None else None,
        "product_score": str(dto["productScore"]) if dto.get("productScore") is not None else None},
        {"merchant_id": merchant_id, "data_date": data_date})
    return {"success": True}

def upload_product_count(db, merchant_id, dto):
    _assert_phase2(db, merchant_id)
    data_date = _normalized_form_data_date(dto.get("dataDate"))
    _upsert(db, ShopProductCount, {"merchant_id": merchant_id, "data_date": data_date, "total_count": dto["totalCount"],
        "on_sale_count": dto.get("onSaleCount"), "off_sale_count": dto.get("offSaleCount"), "audit_count": dto.get("auditCount")},
        {"merchant_id": merchant_id, "data_date": data_date})
    return {"success": True}

def upload_health_score(db, merchant_id, dto):
    _assert_phase2(db, merchant_id)
    data_date = _normalized_form_data_date(dto.get("dataDate"))
    _upsert(db, ShopHealthScore, {"merchant_id": merchant_id, "data_date": data_date, "avg_score": str(dto["avgScore"]),
        "score_gte_90_count": dto.get("scoreGte90Count"), "score_78_90_count": dto.get("score78_90Count"),
        "score_60_77_count": dto.get("score60_77Count"), "score_lt_60_count": dto.get("scoreLt60Count")},
        {"merchant_id": merchant_id, "data_date": data_date})
    return {"success": True}

FIELD_MAPS = {
    "trade": {"成交金额": "trade_amount", "成交单量": "trade_orders", "成交客户数": "trade_customers", "店铺访客数": "shop_visitors",
        "店铺浏览量": "shop_page_views", "成交商品件数": "trade_items", "店铺成交转化率": "conversion_rate", "客单价": "customer_unit_price",
        "店铺平均停留时长": "avg_stay_duration", "加购客户数": "cart_customers", "加购商品件数": "cart_items", "加购转化率": "cart_conversion_rate", "时间": "data_date"},
    "traffic": {"店铺访客数": "shop_visitors", "店铺浏览量": "shop_page_views", "店铺平均停留时长": "avg_stay_duration",
        "商品访客数": "product_visitors", "商品浏览量": "product_page_views", "商品人均浏览量": "product_avg_page_views",
        "商品平均停留时长": "product_avg_stay_duration", "UV价值": "uv_value", "客单价": "customer_unit_price",
        "商品曝光次数": "product_exposure_count", "商品曝光人数": "product_exposure_users", "成交客户数": "trade_customers",
        "加购客户数": "cart_customers", "加购转化率": "cart_conversion_rate", "加购金额": "cart_amount", "成交转化率": "trade_conversion_rate",
        "成交商品件数": "trade_items", "成交单量": "trade_orders", "成交金额": "trade_amount", "时间": "data_date"},
    "product": {"动销spu数": "active_spu_count", "spu动销率": "spu_active_rate", "成交金额": "trade_amount", "成交商品件数": "trade_items",
        "成交单量": "trade_orders", "成交客户数": "trade_customers", "成交转化率": "trade_conversion_rate", "客单价": "customer_unit_price",
        "件单价": "item_unit_price", "加购spu数": "cart_spu_count", "spu加购率": "spu_cart_rate", "加购商品件数": "cart_items",
        "加购客户数": "cart_customers", "加购金额": "cart_amount", "访问spu数": "visit_spu_count", "商品浏览量": "product_page_views",
        "商品访客数": "product_visitors", "商品人均浏览量": "product_avg_page_views", "商详平均停留时长": "product_avg_stay_duration",
        "上架spu数": "listed_spu_count", "时间": "data_date"},
}

# Excel 数值列的取值域(唯一 owner:project/scripts/schema.sql 的真实类型):
#   INT UNSIGNED -> [0, 4294967295];DECIMAL(M,D) -> [0, 10^(M-D) - 10^(-D)]。
# 计数/金额/比率/时长一律**非负**(用户 2026-09-15 裁决:成交单量、成交金额不能为负);
# 越界 -> 跳过该行(不写库、不 500);逐行问题清单与响应契约归 #PB-24-2。
UNSIGNED_INT_MAX = 4294967295
EXCEL_UNSIGNED_INT_COLUMNS: Dict[str, tuple] = {
    "trade": ("trade_orders", "trade_customers", "shop_visitors", "shop_page_views", "trade_items",
              "cart_customers", "cart_items"),
    "traffic": ("shop_visitors", "shop_page_views", "product_visitors", "product_page_views",
                "product_exposure_count", "product_exposure_users", "trade_customers", "cart_customers",
                "trade_items", "trade_orders"),
    "product": ("active_spu_count", "trade_items", "trade_orders", "trade_customers", "cart_spu_count",
                "cart_items", "cart_customers", "visit_spu_count", "product_page_views", "product_visitors",
                "listed_spu_count"),
}
EXCEL_DECIMAL_COLUMNS: Dict[str, Dict[str, tuple]] = {
    "trade": {"trade_amount": (12, 2), "conversion_rate": (5, 2), "customer_unit_price": (10, 2),
              "avg_stay_duration": (8, 2), "cart_conversion_rate": (5, 2)},
    "traffic": {"avg_stay_duration": (8, 2), "product_avg_page_views": (5, 2),
                "product_avg_stay_duration": (8, 2), "uv_value": (10, 2), "customer_unit_price": (10, 2),
                "cart_conversion_rate": (5, 2), "cart_amount": (12, 2), "trade_conversion_rate": (5, 2),
                "trade_amount": (12, 2)},
    "product": {"spu_active_rate": (5, 2), "trade_amount": (12, 2), "trade_conversion_rate": (5, 2),
                "customer_unit_price": (10, 2), "item_unit_price": (10, 2), "spu_cart_rate": (5, 2),
                "cart_amount": (12, 2), "product_avg_page_views": (5, 2),
                "product_avg_stay_duration": (8, 2)},
}


def _is_no_data_marker(value: Any) -> bool:
    """是否为「无数据」标记(商智导出的 - / — / 等):是 -> 该列写 NULL(无值),不是坏值。"""
    return isinstance(value, str) and value.strip() in NO_DATA_MARKERS


def _excel_value_in_domain(file_type: str, column: str, value: Any) -> bool:
    """Excel 数值列的取值域判定(唯一 owner):非负 + 不超列上界;未登记的列不判(返回 True)。"""
    decimal_spec = EXCEL_DECIMAL_COLUMNS[file_type].get(column)
    if decimal_spec is not None:
        precision, scale = decimal_spec
        upper = Decimal(10) ** (precision - scale) - Decimal(1).scaleb(-scale)
    elif column in EXCEL_UNSIGNED_INT_COLUMNS[file_type]:
        upper = Decimal(UNSIGNED_INT_MAX)
    else:
        return True
    number = Decimal(str(value))
    return bool(Decimal(0) <= number <= upper)


# 逐行问题清单的 reason 枚举(唯一真源):日期类复用 DATE_PARSE_REASONS,数值/跨度类在此登记;
# 不新增第三套字面量 —— 路由与服务都从 ISSUE_MESSAGES 取文案。
NUMERIC_ISSUE_REASONS = ("not_a_number", "negative_not_allowed", "out_of_range", "scale_rounded",
                         "invalid_time_range")
# date_normalized = 日期被显式归一(不是错误,action=normalized;如区间取结束日、2026/8/5 -> 2026-08-05)
ISSUE_REASONS = DATE_PARSE_REASONS + ("date_normalized",) + NUMERIC_ISSUE_REASONS

# reason -> 用户可读中文文案(口语化,**不得出现英文字段名**;≤40 字;单一真源)
ISSUE_MESSAGES: Dict[str, str] = {
    "missing_date": "该行缺少数据日期，已跳过",
    "unparseable_date": "日期格式无法识别，已跳过该行",
    "ambiguous_date_format": "日期无法确定年月日顺序，已跳过该行",
    "two_digit_year": "日期用了两位年份，已跳过该行",
    "missing_year": "日期缺少年份，已跳过该行",
    "incomplete_date": "日期不完整，已跳过该行",
    "invalid_calendar_date": "日期不存在，已跳过该行",
    "date_serial_number_not_supported": "日期疑似表格序列号，请设为日期格式后重导",
    "date_out_of_range": "数据日期超出允许范围，已跳过该行",
    "date_normalized": "日期已按规范格式记录",
    "not_a_number": "数值无法识别，已跳过该行",
    "negative_not_allowed": "数值不能为负，已跳过该行",
    "out_of_range": "数值超出允许上限，已跳过该行",
    "scale_rounded": "小数位已按该列精度四舍五入",
    "invalid_time_range": "文件日期跨度不支持，请按单日、7 天或 30 天导出",
}

# 逐行问题清单上界(禁止无界响应):最多回 200 条,超出置 issues_truncated=true;计数仍为全量,完整清单进日志
MAX_ISSUES = 200
# 清单里 `value` 原值回显的最大字符数(截断,避免把整段文本带回客户端)
ISSUE_VALUE_MAX_CHARS = 40

# 文件日期跨度(含首尾天数) -> time_range 桶(用户 2026-09-15 裁决:单日沿用 yesterday)
TIME_RANGE_SINGLE_DAY = "yesterday"
TIME_RANGE_WEEK = "7d"
TIME_RANGE_MONTH = "30d"
SPAN_SINGLE_DAY_MAX = 1
SPAN_WEEK_MAX = 10
SPAN_MONTH_MAX = 31
TIME_RANGE_UNSUPPORTED_MESSAGE = "文件的日期跨度不支持，请按单日、近 7 天或近 30 天导出后重试"
# 表头完全无法识别(一个被映射列都没有):与「数据全坏」区分开,避免把传错文件说成数据问题
HEADER_UNRECOGNIZED_MESSAGE = "文件表头无法识别，请使用商智导出的原始文件"

# 数据看板时间范围枚举(唯一真源):yesterday / 7d / 30d。
# #PB-30:非法值 -> 400 **中文**专句(原 analysis 侧是英文 NestJS 文案,违反真源 §4.3.1「用户可见文案一律中文」);
# 常量与校验函数只此一份,summary 与 analysis 两个入口共用,禁止各拼一句。
TIME_RANGE_VALUES = (TIME_RANGE_SINGLE_DAY, TIME_RANGE_WEEK, TIME_RANGE_MONTH)
TIME_RANGE_INVALID_MESSAGE = "时间范围不支持，请使用 昨天 / 近7天 / 近30天"


def assert_time_range(value: Any) -> str:
    """校验数据看板 `time_range`;非法值抛 400(中文专句,文案单一真源)。

    调用点只有路由边界 2 处(`GET /api/shop/summary`、`POST /api/shop/analysis`),且都**先于**阶段二门禁:
    非法入参恒 400,不会因商家未解锁而变成 403。
    """
    if value not in TIME_RANGE_VALUES:
        raise ApiException(TIME_RANGE_INVALID_MESSAGE, code=400, status_code=400)
    return value


def _column_scale(file_type: str, column: str) -> Optional[int]:
    """该列的精度 scale:INT UNSIGNED -> 0;DECIMAL(M,D) -> D;未登记 -> None(不量化)。"""
    decimal_spec = EXCEL_DECIMAL_COLUMNS[file_type].get(column)
    if decimal_spec is not None:
        return decimal_spec[1]
    if column in EXCEL_UNSIGNED_INT_COLUMNS[file_type]:
        return 0
    return None


def _quantize_to_scale(value: Any, scale: int) -> Any:
    """按列精度**显式**量化(ROUND_HALF_UP,与 MySQL 取整口径一致),不再交给 MySQL 隐式四舍五入。"""
    quantized = Decimal(str(value)).quantize(Decimal(1).scaleb(-scale), rounding=ROUND_HALF_UP)
    return int(quantized) if scale == 0 else float(quantized)


def time_range_from_span(start: date, end: date) -> str:
    """文件日期跨度(含首尾) -> `time_range` 桶:1 天 -> yesterday;2-10 天 -> 7d;11-31 天 -> 30d。

    其它跨度(含起晚于止)-> 400(TIME_RANGE_UNSUPPORTED_MESSAGE,整份文件拒绝):跨度对不上任何看板窗口时,
    与其塞进某个桶让看板显示错数字,不如让用户按受支持的跨度重新导出。
    """
    days = (end - start).days + 1
    if days <= SPAN_SINGLE_DAY_MAX or days <= 0:
        return TIME_RANGE_SINGLE_DAY if days == SPAN_SINGLE_DAY_MAX else _raise_unsupported_span(start, end)
    if days <= SPAN_WEEK_MAX:
        return TIME_RANGE_WEEK
    if days <= SPAN_MONTH_MAX:
        return TIME_RANGE_MONTH
    return _raise_unsupported_span(start, end)


def _raise_unsupported_span(start: date, end: date) -> str:
    """不支持的跨度:结构化 400(口语化文案;reason 记为 invalid_time_range,只进日志)。"""
    days = (end - start).days + 1
    logger.warning(f"upload_excel 拒绝导入: 文件日期跨度不支持 (reason=invalid_time_range, days={days})")
    raise ApiException(TIME_RANGE_UNSUPPORTED_MESSAGE, code=400, status_code=400)


def _issue_value(value: Any) -> str:
    """清单里的原值回显:截断到 ISSUE_VALUE_MAX_CHARS(只取日期/数值列,不含其它列内容)。"""
    return ("" if value is None else str(value))[:ISSUE_VALUE_MAX_CHARS]


def _append_issue(issues: List[Dict[str, Any]], row: int, field: str, value: Any,
                  reason: str, action: str = "skipped") -> None:
    """登记一条逐行问题(文案取自 ISSUE_MESSAGES 单一真源,调用点不得另拼文案)。"""
    issues.append({"row": row, "field": field, "value": _issue_value(value),
                   "reason": reason, "action": action, "message": ISSUE_MESSAGES[reason]})


def _log_issues(file_type: str, issues: List[Dict[str, Any]]) -> None:
    """把**完整**清单写日志(有界:前 MAX_ISSUES 条逐条 + 其余条数汇总)。"""
    if not issues:
        return
    for item in issues[:MAX_ISSUES]:
        logger.warning("upload_excel 问题清单[" + file_type + "]: " + json.dumps(item, ensure_ascii=False))
    if len(issues) > MAX_ISSUES:
        logger.warning(
            f"upload_excel 问题清单[{file_type}]: 其余 {len(issues) - MAX_ISSUES} 条未逐条打印(计数已全量返回)"
        )

def _normalized_form_data_date(value: Any) -> str:
    """表单路径的日期归一(唯一 owner = parse_data_date):失败 -> 400。

    文案取自模块常量 DATA_DATE_INVALID_MESSAGE(**单一真源**,不拼接通用兜底句);
    不新增自定义错误码(真源 §9.4)。
    """
    try:
        parsed = parse_data_date(value)
    except DataDateParseError:
        raise ApiException(DATA_DATE_INVALID_MESSAGE, code=400, status_code=400)
    if not parsed:   # 理论上不可达(无值由 parse_data_date 抛 missing_date),防御性保留同一 400
        raise ApiException(DATA_DATE_INVALID_MESSAGE, code=400, status_code=400)
    return parsed

def _to_num(v):
    """数值转换(Excel 导入策略):脏值跳过(记日志返回 None),不抛断整批。
    解析规则归 parse_number(数值解析唯一 owner),本函数只决定"失败即跳过"。
    整数值归一为 int(与 NestJS toNumber 的 number 语义一致,避免 5.0 这类浮点形态写进 INT 列)。"""
    try:
        num = parse_number(v)
    except NumberParseError:
        logger.warning(f"_to_num 跳过无法解析的脏值: {v!r}")
        return None
    if num is None:
        return None
    if num.is_finite() and num == num.to_integral_value():
        return int(num)
    return float(num)

def upload_excel(db, merchant_id, file_type: str, raw: bytes, mime: str) -> Dict[str, Any]:
    """Excel 导入(行级):坏行跳过 + **逐行问题清单**;`time_range` 按文件日期跨度推导(#PB-24-2)。

    失败策略(用户 2026-09-15 裁决,与方案 §8 一致):
    - 表头完全无法识别 / 文件日期跨度不支持 / 无数据行 -> **400**(不写库、文案口语化);
    - 单行内任一被映射的日期/数值单元格坏(含日期缺失)-> **跳过该行**并进清单(action=skipped);
    - `-` 等无数据标记 -> 该列写 NULL(**不跳行**、不进清单);数值小数位超列精度 -> 显式四舍五入并进清单
      (action=normalized);日期被显式归一(区间取结束日 / 非规范形态)-> 进清单(action=normalized);
    - 全部行皆坏 -> **200 + 清单**(不再 400 丢清单);任何输入都不得 500。
    - 清单上界 MAX_ISSUES,超出置 issues_truncated 且 count/total_rows/skipped/normalized 仍为全量;完整清单进日志。
    """
    _assert_phase2(db, merchant_id)
    if mime not in EXCEL_MIME:
        raise ApiException("仅支持 .xlsx 文件，.xls 请另存为 .xlsx 后上传", code=400, status_code=400)
    try:
        wb = load_workbook(io.BytesIO(raw), data_only=True)
    except Exception:
        raise ApiException("Excel 文件解析失败，请上传正确的 Excel 文件", code=400, status_code=400)
    if not wb.sheetnames:
        raise ApiException("Excel 中未解析到有效数据", code=400, status_code=400)
    if len(wb.sheetnames) > MAX_SHEETS:
        raise ApiException("Excel sheet 数量过多，最多支持 10 个", code=400, status_code=400)
    ws = wb[wb.sheetnames[0]]
    if ws.max_row > MAX_ROWS:
        raise ApiException("Excel 数据量过大，最多支持 5000 行", code=400, status_code=400)
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ApiException("Excel 中未解析到有效数据", code=400, status_code=400)
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    maps = FIELD_MAPS[file_type]
    if not any(h in maps for h in header):
        # 一个被映射列都没有:大概率传错文件/表头被改(与「数据内容有坏行」区分开)
        logger.warning(f"upload_excel 表头无法识别: file_type={file_type} header={header[:10]}")
        raise ApiException(HEADER_UNRECOGNIZED_MESSAGE, code=400, status_code=400)
    unmapped = [h for h in header if h and h not in maps]
    if unmapped:
        # 用户裁决「未导入的列**不要提示用户**」-> 只写日志(有界),不进响应契约
        logger.warning(
            f"upload_excel 未映射的列(不导入,仅日志): file_type={file_type} count={len(unmapped)} "
            f"headers={unmapped[:50]}"
        )
    issues: List[Dict[str, Any]] = []
    count = 0
    total_rows = 0
    skipped = 0
    normalized = 0
    time_range: Optional[str] = None
    model = {"trade": ShopTradeData, "traffic": ShopTrafficData, "product": ShopProductData}[file_type]
    try:
        for excel_row, r in enumerate(rows[1:], start=2):   # excel_row = Excel 物理行号(表头为第 1 行)
            if r is None or all(c is None for c in r):
                continue
            total_rows += 1
            rec: Dict[str, Any] = {}
            row_rejected = False
            row_normalized = False
            for i, col in enumerate(header):
                en_key = maps.get(col)
                if not en_key:
                    continue
                val = r[i] if i < len(r) else None
                if val is None or val == "":
                    continue
                if en_key == "data_date":
                    # 日期解析归 data_date_span(唯一 owner,复用区间跨度,不重复解析)。
                    # 失败策略 = 跳过该行(绝不兜底成「今天」:旧实现会把空白日期静默写成今天并可能覆盖当天数据)。
                    try:
                        start, end = data_date_span(val)
                    except DataDateParseError as exc:
                        _append_issue(issues, excel_row, en_key, val, exc.reason)
                        logger.warning(f"upload_excel 第 {excel_row} 行数据日期不可用: reason={exc.reason}")
                        row_rejected = True
                        break
                    if time_range is None:
                        # 文件级跨度:取**首个可解析日期行**的跨度(同一份导出内跨度一致)
                        time_range = time_range_from_span(start, end)
                    canonical = end.isoformat()
                    rec[en_key] = canonical
                    if not isinstance(val, (date, datetime)) and str(val).strip() != canonical:
                        _append_issue(issues, excel_row, en_key, val, "date_normalized", action="normalized")
                        row_normalized = True
                    continue
                if _is_no_data_marker(val):
                    continue          # 无数据标记 -> 该列写 NULL,不跳过整行、不进清单
                num = _to_num(val)
                if num is None:
                    _append_issue(issues, excel_row, en_key, val, "not_a_number")
                    row_rejected = True
                    break
                if num < 0:
                    _append_issue(issues, excel_row, en_key, val, "negative_not_allowed")
                    row_rejected = True
                    break
                if not _excel_value_in_domain(file_type, en_key, num):
                    _append_issue(issues, excel_row, en_key, val, "out_of_range")
                    row_rejected = True
                    break
                scale = _column_scale(file_type, en_key)
                if scale is not None:
                    quantized = _quantize_to_scale(num, scale)
                    if quantized != num:
                        _append_issue(issues, excel_row, en_key, val, "scale_rounded", action="normalized")
                        row_normalized = True
                    num = quantized
                rec[en_key] = num
            if row_rejected:
                skipped += 1
                continue
            if not rec:
                continue
            if not rec.get("data_date"):
                # 日期列整列未映射(表头被改成未知形态)等:同一套 missing_date 处理,不得再抛 KeyError/500
                _append_issue(issues, excel_row, "data_date", None, "missing_date")
                skipped += 1
                continue
            data_date = rec.pop("data_date")
            rec["merchant_id"] = merchant_id; rec["data_date"] = data_date; rec["time_range"] = time_range
            # 整批一个事务:循环内不 commit,循环外统一 commit
            _upsert(db, model, rec, {"merchant_id": merchant_id, "data_date": data_date, "time_range": time_range}, commit=False)
            count += 1
            if row_normalized:
                normalized += 1
        if total_rows == 0:
            raise ApiException("Excel 中未解析到有效数据", code=400, status_code=400)
        _log_issues(file_type, issues)
        db.commit()
        return {"success": True, "count": count, "total_rows": total_rows, "skipped": skipped,
                "normalized": normalized, "issues": issues[:MAX_ISSUES],
                "issues_truncated": len(issues) > MAX_ISSUES}
    except Exception:
        db.rollback()
        raise


def get_summary(db, merchant_id, time_range):
    _assert_phase2(db, merchant_id)
    def latest(model, where):
        return db.execute(select(model).where(*[getattr(model, k) == v for k, v in where.items()]).order_by(model.data_date.desc()).limit(1)).scalars().first()
    star = latest(ShopStarData, {"merchant_id": merchant_id})
    trade = latest(ShopTradeData, {"merchant_id": merchant_id, "time_range": time_range})
    traffic = latest(ShopTrafficData, {"merchant_id": merchant_id, "time_range": time_range})
    product = latest(ShopProductData, {"merchant_id": merchant_id, "time_range": time_range})
    pc = latest(ShopProductCount, {"merchant_id": merchant_id})
    hs = latest(ShopHealthScore, {"merchant_id": merchant_id})
    return {"time_range": time_range,
            "star": _serialize_decimals({"shop_star": star.shop_star, "service_score": star.service_score, "logistics_score": star.logistics_score, "after_sale_score": star.after_sale_score, "product_score": star.product_score, "data_date": star.data_date}, "star") if star else None,
            "trade": _serialize_decimals({"trade_amount": trade.trade_amount, "trade_orders": trade.trade_orders, "trade_customers": trade.trade_customers, "shop_visitors": trade.shop_visitors, "shop_page_views": trade.shop_page_views, "trade_items": trade.trade_items, "conversion_rate": trade.conversion_rate, "customer_unit_price": trade.customer_unit_price, "avg_stay_duration": trade.avg_stay_duration, "cart_customers": trade.cart_customers, "cart_items": trade.cart_items, "cart_conversion_rate": trade.cart_conversion_rate, "data_date": trade.data_date, "time_range": trade.time_range}, "trade") if trade else None,
            "traffic": _serialize_decimals({c: getattr(traffic, c) for c in ("shop_visitors", "shop_page_views", "avg_stay_duration", "product_visitors", "product_page_views", "product_avg_page_views", "product_avg_stay_duration", "uv_value", "customer_unit_price", "product_exposure_count", "product_exposure_users", "trade_customers", "cart_customers", "cart_conversion_rate", "cart_amount", "trade_conversion_rate", "trade_items", "trade_orders", "trade_amount")} | {"data_date": traffic.data_date, "time_range": traffic.time_range}, "traffic") if traffic else None,
            "product": _serialize_decimals({c: getattr(product, c) for c in ("active_spu_count", "spu_active_rate", "trade_amount", "trade_items", "trade_orders", "trade_customers", "trade_conversion_rate", "customer_unit_price", "item_unit_price", "cart_spu_count", "spu_cart_rate", "cart_items", "cart_customers", "cart_amount", "visit_spu_count", "product_page_views", "product_visitors", "product_avg_page_views", "product_avg_stay_duration", "listed_spu_count")} | {"data_date": product.data_date, "time_range": product.time_range}, "product") if product else None,
            "product_count": {"total_count": pc.total_count, "on_sale_count": pc.on_sale_count, "off_sale_count": pc.off_sale_count, "audit_count": pc.audit_count, "data_date": pc.data_date} if pc else None,
            "health_score": _serialize_decimals({"avg_score": hs.avg_score, "score_gte_90_count": hs.score_gte_90_count, "score_78_90_count": hs.score_78_90_count, "score_60_77_count": hs.score_60_77_count, "score_lt_60_count": hs.score_lt_60_count, "data_date": hs.data_date}, "health_score") if hs else None}


def clear_merchant_data(db: Session, merchant_id: str) -> Dict[str, Any]:
    """清除当前所有数据(破坏性,范围已锁定):
    ① 删除该 merchant 在 6 张 shop_* 表记录;
    ② 保留 ai_analysis_log(不清);
    ③ 仅重置 merchant_task_progress 中 T2.5.x 完成态为 pending;
    ④ 保留数据专区解锁(不改 merchant.data_center_unlocked、不清 merchant_stage_progress);
    ⑤ 整个清除在一个事务内,异常回滚。
    """
    if not isinstance(merchant_id, str) or not merchant_id.strip():
        raise ApiException("merchant_id 非法", code=400, status_code=400)
    table_map = {
        "star": "shop_star_data",
        "trade": "shop_trade_data",
        "traffic": "shop_traffic_data",
        "product": "shop_product_data",
        "product_count": "shop_product_count",
        "health_score": "shop_health_score",
    }
    try:
        cleared = {}
        for key, table in table_map.items():
            res = db.execute(text(f"DELETE FROM {table} WHERE merchant_id = :m"), {"m": merchant_id})
            cleared[key] = int(res.rowcount or 0)
        res2 = db.execute(text(
            "UPDATE merchant_task_progress SET status = 'pending', completedAt = NULL "
            "WHERE merchantId = :m AND taskId LIKE 'T2.5.%' AND status = 'completed'"
        ), {"m": merchant_id})
        reset_tasks = int(res2.rowcount or 0)
        db.commit()
        return {"success": True, "cleared": cleared, "reset_tasks": reset_tasks}
    except Exception as e:
        db.rollback()
        # SEC P0:清除数据失败(不该吞)记日志后上抛,全局异常返回明确失败
        logger.error(f"清除数据失败: merchant={merchant_id} err={e}")
        raise